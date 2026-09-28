"""Write a Desktop Word note: Pro Forma understanding + red unlocked items."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

out_dir = Path(r"C:\Users\Administrator\Desktop\LexZim Proforma Discussion")
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / "Proforma Invoice Understanding.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)


def add_para(text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    run.font.color.rgb = BLACK
    return p


def add_mixed(parts):
    p = doc.add_paragraph()
    for text, color, bold in parts:
        run = p.add_run(text)
        run.font.color.rgb = color
        run.bold = bold
        run.font.size = Pt(11)
    return p


def bullet(text: str, color=BLACK, bold: bool = False):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.color.rgb = color
    run.bold = bold
    run.font.size = Pt(11)
    return p


def bullet_mixed(parts):
    p = doc.add_paragraph(style="List Bullet")
    for text, color, bold in parts:
        run = p.add_run(text)
        run.font.color.rgb = color
        run.bold = bold
        run.font.size = Pt(11)
    return p


doc.add_heading("LexZim Pro Forma Invoice", 0)
add_para("What we understand so far (simple English)")
add_para("Firm scope: T Pfigu Attorneys only for this next update.")
add_para("Date of this note: 10 September 2026")
add_mixed(
    [
        ("Legend: ", BLACK, True),
        ("Black text = agreed / locked. ", BLACK, False),
        (
            "Red text = not locked yet, or still confused. Needs a clear yes/no before we build.",
            RED,
            True,
        ),
    ]
)

doc.add_heading("1. Big picture", level=1)
bullet("Keep the normal general invoice / fee note as it is today.")
bullet("Add a separate document type: Pro Forma Invoice for property transfer and related work.")
bullet("A Pro Forma is an estimate. It can change. That is why it is called a Pro Forma.")
bullet(
    "It should not post to debtors or the general ledger until it is later turned into a real invoice or trust request."
)
bullet("All Pro Forma bank details go to the Trust account (not the firm operating account).")
bullet("Sample files used for discussion were Grandia (transfer / developer style) and Zokwa (cession).")

doc.add_heading("2. Types of Pro Forma", level=1)
bullet("Normal property transfer")
bullet("Developer property (price may show VAT)")
bullet("Cession")
add_mixed(
    [
        (
            "We still need to confirm the exact screen labels and whether these are three separate templates or one form with a type choice.",
            RED,
            True,
        )
    ]
)

doc.add_heading("3. What appears on a Pro Forma (typical lines)", level=1)

doc.add_heading("3.1 Normal transfer", level=2)
bullet("Stamp duty (calculated from brackets)")
bullet("Registration fee: always USD 50")
bullet(
    "Conveyancing fee: default 3% of the relevant value, but staff can change it (negotiation allowed)"
)
bullet("Disbursements: typed in by staff (amounts vary; no fixed system default)")
bullet("Admin / other fees: typed in by staff (manual)")
bullet_mixed(
    [
        (
            "Any other standard lines (for example rates clearances wording) are not fully listed yet.",
            RED,
            True,
        )
    ]
)

doc.add_heading("3.2 Developer property", level=2)
bullet("Can show a display price that includes VAT")
bullet("Needs a separate dutiable value field used only for stamp duty (two fields)")
bullet(
    "Developer-related fees vary by service provider / council: typed in, not fixed by the system"
)
bullet_mixed(
    [
        (
            'Exact VAT wording on the PDF (for example "VAT paid by purchaser") still needs a final sample sentence.',
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "Whether Capital Gains Tax is only a note, or also an amount line, is not locked.",
            RED,
            True,
        )
    ]
)

doc.add_heading("3.3 Cession", level=2)
bullet("No stamp duty")
bullet("Registration fee USD 50: never")
bullet("Cession / council / attendance fees: vary by council or service provider; typed in")
bullet("Conveyancing can still appear, often as a negotiated amount")
bullet("Disbursements: typed in")
bullet_mixed(
    [
        (
            "Default suggested amounts for cession lines (Zokwa had 500 / 250 / 500 / 500) are not locked. Treat as examples only for now.",
            RED,
            True,
        )
    ]
)

doc.add_heading("4. Stamp duty brackets (main remaining discussion)", level=1)
add_para("What we understand:")
bullet("Stamp duty uses brackets. The property value falls into one bracket.")
bullet("Each bracket has a percentage and a flat fee.")
add_para("Your latest explanation (as understood):")
bullet("In the 2% bracket: 2% of the actual (full) property value, plus flat fee USD 70.")
bullet("In the 3% bracket: 3% of the actual (full) property value, plus that bracket flat fee (USD 370).")
bullet(
    "Same idea for other brackets: percentage of the actual value in that deal, plus the flat fee for that bracket."
)
add_mixed(
    [
        (
            'Earlier examples sometimes used "excess over the previous band" '
            "(for example 50,000 as (50,000 - 20,000) x 3% + 370). "
            'That conflicts with "percentage of the full actual value". '
            "This is the biggest open confusion.",
            RED,
            True,
        )
    ]
)

doc.add_heading("4.1 Draft bracket table (partly agreed, partly red)", level=2)

table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Property value band"
hdr[1].text = "What we think the rule is"
hdr[2].text = "Status"

rows_data = [
    (
        "USD 0 to 5,000",
        "1% of full value + USD 70\nExample: 4,000 -> 110",
        "Mostly agreed",
    ),
    (
        "Above 5,000 to 20,000",
        "2% of full value + USD 70\nExample: 10,000 -> 270",
        "Mostly agreed (if full-value rule stands)",
    ),
    (
        "Above 20,000 to 100,000",
        "Either:\nA) 3% of full value + 370\n   50,000 -> 1,870\n"
        "Or B) 3% of (value - 20,000) + 370\n   50,000 -> 1,270",
        "NOT LOCKED",
    ),
    (
        "Above 100,000",
        "Either:\nA) 4% of full value + 2,770\n   500,000 -> 22,770\n"
        "Or B) 4% of (value - 100,000) + 2,770\n   500,000 -> 18,770",
        "NOT LOCKED",
    ),
]
for i, (a, b, c) in enumerate(rows_data, start=1):
    table.rows[i].cells[0].text = a
    table.rows[i].cells[1].text = b
    table.rows[i].cells[2].text = c

for i in (3, 4):
    cell = table.rows[i].cells[2]
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run("NOT LOCKED / CONFUSION")
    run.font.color.rgb = RED
    run.bold = True

add_para("")
add_mixed([("Flat fee for the mid band is locked as USD 370 (not 3,700).", BLACK, False)])
add_mixed(
    [
        (
            'Whether there is any 5% bracket: you mentioned "5% or 3%" in speech. '
            "We currently assume only 1%, 2%, 3%, and 4%. Please confirm there is no 5% bracket.",
            RED,
            True,
        )
    ]
)
add_mixed(
    [
        (
            'For 10,000, one message said stamp duty "70". That may have meant only the flat fee, '
            "not the total. If total is only 70, that conflicts with 2% + 70. "
            "Needs clear confirmation of the total for 10,000.",
            RED,
            True,
        )
    ]
)

doc.add_heading("4.2 Developer stamp duty base", level=2)
bullet("Do not automatically divide the display price by 1.155 in the system.")
bullet("Use two fields:")
bullet("Field 1: Display price (may include VAT)")
bullet("Field 2: Dutiable value for stamp duty")
bullet("Stamp duty calculator uses Field 2 only.")
add_mixed(
    [
        (
            "Grandia only matched the printed stamp duty when duty was calculated on about "
            "547,086 (VAT stripped), not on the gross 631,884.33. With two fields, staff would "
            "enter that dutiable amount in Field 2. Confirm that this is how T Pfigu works in practice.",
            RED,
            True,
        )
    ]
)

doc.add_heading("5. Conveyancing, fees, and money lines", level=1)
bullet("Conveyancing default: 3%. Staff may override (negotiation). LOCKED.")
bullet("Admin fee: manual. LOCKED.")
bullet("Disbursements: vary; staff type the amount. No fixed default. LOCKED.")
bullet("Registration fee on normal / developer transfer: always USD 50. LOCKED.")
bullet("Registration fee on cession: never. LOCKED.")
bullet("Developer fee / cession fee: varies by provider or council; manual. LOCKED.")
bullet("All Pro Forma payments directed to Trust banking details. LOCKED.")
add_mixed(
    [
        (
            "Grandia sample showed account name without the word Trust; Zokwa showed Trust. "
            "Going forward we treat Trust as correct for all Pro Formas. "
            "Old Grandia wording was a sample inconsistency, not the rule.",
            BLACK,
            False,
        )
    ]
)

doc.add_heading("6. Accounting behaviour", level=1)
bullet("Pro Forma is subject to change.")
bullet(
    "Recommended build: document / estimate only. No debtors balance and no GL posting until converted later."
)
add_mixed(
    [
        (
            "Exact convert path later (Pro Forma to trust request, fee invoice, or both) is not designed yet.",
            RED,
            True,
        )
    ]
)

doc.add_heading("7. What LexZim has today", level=1)
bullet("One generic invoice / fee note.")
bullet("No Pro Forma subtype.")
bullet("No stamp duty bracket calculator.")
bullet("Quotations exist, but they are not this property Pro Forma.")

doc.add_heading("8. Open questions checklist (all red until answered)", level=1)
bullet_mixed(
    [
        (
            "For 50,000 normal transfer, is stamp duty 1,870 (full x 3% + 370) or 1,270 (excess x 3% + 370)?",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "For 500,000, is stamp duty 22,770 or 18,770?",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "Is there a 5% stamp duty bracket, or only 1% / 2% / 3% / 4%?",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "For 10,000, confirm the TOTAL stamp duty number T Pfigu would put on a Pro Forma.",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "Confirm two fields for developer: display price + dutiable value, and stamp duty uses dutiable value only.",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "Confirm default cession line suggestions, if any, or always blank for staff to type.",
            RED,
            True,
        )
    ]
)
bullet_mixed(
    [
        (
            "Confirm later conversion from Pro Forma into trust / invoice (even if build comes later).",
            RED,
            True,
        )
    ]
)

doc.add_heading("9. Bottom line", level=1)
add_para(
    "We understand the product goal: a T Pfigu-only Pro Forma for property work, "
    "separate from the general invoice, with Trust banking, estimate-only behaviour, "
    "registration rules, and mostly manual fee lines."
)
add_mixed(
    [
        (
            "We cannot safely code the stamp duty engine until the full-value vs excess rule "
            "is locked for the 3% and 4% brackets, and the 10,000 total is confirmed.",
            RED,
            True,
        )
    ]
)
add_para(
    "Developer two-field approach (display price + dutiable value) is understood and "
    "ready to discuss next once stamp duty brackets are clear."
)
add_para("")
add_para("End of note.", bold=True)

doc.save(str(out_path))
print(out_path)
