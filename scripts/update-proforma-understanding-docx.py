"""Pro Forma understanding doc: locked body + green questions at the bottom."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

path = Path(
    r"C:\Users\Administrator\Desktop\LexZim Proforma Discussion"
    r"\Proforma Invoice Understanding (3).docx"
)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x80, 0x40)
GREY = RGBColor(0x55, 0x55, 0x55)


def para(text: str, bold: bool = False, color=BLACK, size: int = 11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def mixed(parts):
    p = doc.add_paragraph()
    for text, color, bold in parts:
        r = p.add_run(text)
        r.font.color.rgb = color
        r.bold = bold
        r.font.size = Pt(11)
    return p


def bullet(text: str, color=BLACK, bold: bool = False):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.color.rgb = color
    r.bold = bold
    r.font.size = Pt(11)
    return p


def green_para(text: str, bold: bool = False, size: int = 11):
    return para(text, bold=bold, color=GREEN, size=size)


def green_bullet(text: str, bold: bool = False):
    return bullet(text, color=GREEN, bold=bold)


def blank_line(label: str = "Others / further explanation:", color=GREEN):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    r.font.color.rgb = color
    r.italic = True
    r.font.size = Pt(10)
    r2 = p.add_run("_" * 72)
    r2.font.color.rgb = color
    r2.font.size = Pt(10)
    for _ in range(2):
        p2 = doc.add_paragraph()
        r3 = p2.add_run("_" * 90)
        r3.font.color.rgb = color
        r3.font.size = Pt(10)
    return p


def ask_green(number: str, title: str, detail: str, choices: str | None = None):
    p = doc.add_paragraph()
    r = p.add_run(f"Question {number}. {title}")
    r.bold = True
    r.font.color.rgb = GREEN
    r.font.size = Pt(11)

    p2 = doc.add_paragraph()
    r2 = p2.add_run(detail)
    r2.font.color.rgb = GREEN
    r2.font.size = Pt(11)

    if choices:
        p3 = doc.add_paragraph()
        r3 = p3.add_run("Please answer: " + choices)
        r3.bold = True
        r3.font.color.rgb = GREEN
        r3.font.size = Pt(11)

    blank_line(color=GREEN)
    para("")


# ---------------------------------------------------------------------------
# Main body: understanding so far (black). Keep structure simple.
# ---------------------------------------------------------------------------

doc.add_heading("LexZim Pro Forma Invoice", 0)
para("What we understand so far (simple English)")
para("Firm scope: T Pfigu Attorneys only for this next update.")
para("Updated: 17 September 2026")
mixed(
    [
        ("Legend: ", BLACK, True),
        ("Black = agreed / locked understanding. ", BLACK, False),
        (
            "Green (at the very bottom) = new questions and everything still needed from you. "
            "Please answer the green section.",
            GREEN,
            True,
        ),
    ]
)

doc.add_heading("1. Big picture (locked)", level=1)
bullet("Keep the normal general invoice / fee note as it is today.")
bullet(
    "Add a separate document type: Pro Forma Invoice for property transfer and related work."
)
bullet(
    "A Pro Forma is an estimate. It can change up or down. That is why it is called a Pro Forma."
)
bullet(
    "It should not post to debtors or the general ledger until it is later turned into a real invoice or trust request."
)
bullet(
    "All Pro Forma bank details go to the Trust account. Put the word Trust on the account name for consistency."
)
bullet("Sample files used for discussion: Grandia (transfer / developer) and Zokwa (cession).")
bullet(
    "Note from T Pfigu: conveyancing (about 3%, or negotiated) is firm income. "
    "Money-flow detail is asked in green at the bottom."
)

doc.add_heading("2. Types of Pro Forma discussed", level=1)
bullet("Normal property transfer")
bullet("Developer property")
bullet("Industrial property (often has VAT; some are exempt)")
bullet("Cession")
para(
    "Screen labels and whether this is one form or separate templates are asked in green at the bottom."
)

doc.add_heading("3. What appears on a Pro Forma", level=1)

doc.add_heading("3.1 Normal transfer (mostly locked)", level=2)
bullet("Stamp duty: calculated from brackets (see section 4).")
bullet("Registration fee: always USD 50. LOCKED.")
bullet(
    "Conveyancing fee: default 3% of the relevant value; staff can change it "
    "(negotiation allowed). LOCKED."
)
bullet("Disbursements: staff type the amount. LOCKED.")
bullet("Admin / other fees: staff type the amount. LOCKED.")
para("Any extra always-on lines are asked in green at the bottom.")

doc.add_heading("3.2 Developer property (mostly locked)", level=2)
bullet(
    "Two value fields: Display price, and Dutiable value for stamp duty. "
    "Stamp duty uses dutiable value only. Do not auto-divide by 1.155. LOCKED."
)
bullet(
    "Do not calculate conveyancing on the VAT-inclusive gross. "
    "Staff may use a negotiated % (example: 1.5% instead of 3%). LOCKED."
)
bullet("Developer-related fees: staff type them (vary by provider / council). LOCKED.")
para("Final VAT wording and Capital Gains Tax treatment are asked in green at the bottom.")

doc.add_heading("3.3 Industrial property", level=2)
bullet("Discussed as usually having VAT; some deals are exempt.")
para("Whether Industrial matches Developer rules is asked in green at the bottom.")

doc.add_heading("3.4 Cession (mostly locked)", level=2)
bullet("No stamp duty. LOCKED.")
bullet("Registration fee USD 50: never. LOCKED.")
bullet(
    "Cession / council / attendance fees: staff type them. "
    "Suggested sample amounts left blank. LOCKED."
)
bullet("Conveyancing can still appear as a negotiated amount. LOCKED.")
bullet("Disbursements: staff type them. LOCKED.")

doc.add_heading("4. Stamp duty brackets (understanding so far)", level=1)
para("What we believe you locked in version (3):", bold=True)
bullet("There is no 5% bracket. Only 1%, 2%, 3%, and 4%. LOCKED.")
bullet("Flat fee in the mid band is USD 370 (not 3,700). LOCKED.")
bullet(
    'When someone said "70" for a USD 10,000 deal, that meant the flat fee only, '
    "not the total. LOCKED."
)
bullet(
    "For the higher bands, you left formula B (excess over the previous band), "
    "with examples: 50,000 -> 1,270 and 500,000 -> 18,770."
)

para("Draft rule we will code IF you confirm yes in the green section below:", bold=True)
table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Property value band"
hdr[1].text = "Proposed rule"
hdr[2].text = "Example"
rows = [
    ("USD 0 to 5,000", "1% of full value + USD 70", "4,000 -> 110"),
    ("Above 5,000 to 20,000", "2% of full value + USD 70", "10,000 -> 270"),
    (
        "Above 20,000 to 100,000",
        "3% of (value - 20,000) + USD 370",
        "50,000 -> 1,270",
    ),
    (
        "Above 100,000",
        "4% of (value - 100,000) + USD 2,770",
        "500,000 -> 18,770",
    ),
]
for i, (a, b, c) in enumerate(rows, start=1):
    table.rows[i].cells[0].text = a
    table.rows[i].cells[1].text = b
    table.rows[i].cells[2].text = c

para("")
para(
    "Please confirm this table in the green questions at the very bottom of this document."
)

doc.add_heading("5. Conveyancing, fees, and money (locked parts)", level=1)
bullet("Conveyancing default 3%; staff may override. LOCKED.")
bullet("Admin, disbursements, developer/cession fees: manual. LOCKED.")
bullet("Registration USD 50 on normal/developer; never on cession. LOCKED.")
bullet("Pro Forma banking details: Trust account, with Trust in the account name. LOCKED.")
para(
    "How conveyancing as firm income fits with Trust payment instructions "
    "is asked in green at the bottom."
)

doc.add_heading("6. Accounting behaviour (locked for first build)", level=1)
bullet("Pro Forma is subject to change (more or less). LOCKED.")
bullet(
    "Recommended first build: estimate document only. "
    "No debtors balance and no GL posting until converted later. LOCKED for phase 1."
)
para("Later conversion path is asked in green at the bottom (can wait for first build).")

doc.add_heading("7. What LexZim has today", level=1)
bullet("One generic invoice / fee note.")
bullet("No Pro Forma subtype.")
bullet("No stamp duty bracket calculator.")
bullet("Quotations exist, but they are not this property Pro Forma.")

doc.add_heading("8. Bottom line", level=1)
para(
    "We understand the goal: a T Pfigu-only property Pro Forma, separate from the "
    "general invoice, Trust banking, estimate-only behaviour, registration rules, "
    "and mostly manual fee lines."
)
para(
    "Everything still needed from you is in GREEN below. Please fill that section."
)

# ---------------------------------------------------------------------------
# VERY BOTTOM: all new questions and open needs in GREEN
# ---------------------------------------------------------------------------

doc.add_page_break()

h = doc.add_heading("PLEASE ANSWER THIS SECTION (NEW / STILL NEEDED)", level=1)
for run in h.runs:
    run.font.color.rgb = GREEN

green_para(
    "Everything in this section is new or still needed before we can safely build. "
    "All text here is green on purpose so it is easy to find. "
    "Under every question there is an Others / further explanation area "
    "so you can write more if the short answer is not enough.",
    bold=True,
)
green_para("")

green_para("A. Screen and Pro Forma types", bold=True)
ask_green(
    "1",
    "How should staff choose the Pro Forma type on screen?",
    "We need to know how the screen should work when creating a Pro Forma. "
    "Option A: one Pro Forma screen with a Type dropdown "
    "(Normal / Developer / Industrial / Cession). "
    "Option B: four separate templates or menu items, one for each type. "
    "Option C: something else (for example two screens only: Transfer and Cession).",
    "A / B / C. If C, describe it on the Others lines.",
)
ask_green(
    "2",
    "What exact names should staff see for each type?",
    "Please write the exact labels you want on the screen. "
    "For example: “Normal Transfer”, “Developer Property”, “Industrial Property”, “Cession”. "
    "If Industrial should be called something else (for example “Commercial / Industrial”), write that.",
    "Write the four labels (or more/less if needed).",
)

green_para("B. Lines and wording on the Pro Forma", bold=True)
ask_green(
    "3",
    "Are there other standard money lines that must always appear on a normal transfer Pro Forma?",
    "Besides stamp duty, registration fee (USD 50), conveyancing, disbursements, and admin, "
    "do you always need other lines? Examples people sometimes use: rates clearance, "
    "search fees, postage, attendance at Deeds Office, municipality fees. "
    "If yes, list every line that should always show (even if the amount starts at zero). "
    "If no, write none.",
    'none / list the line names',
)
ask_green(
    "4",
    "What exact VAT sentence must appear on the Developer Pro Forma PDF?",
    "Please paste the final wording you want printed. "
    "For example: “VAT paid by purchaser” or a longer sentence. "
    "If VAT wording changes by deal, say whether staff type it each time, "
    "or choose from a short list of standard sentences.",
    "Paste final sentence, or say staff type it each time.",
)
ask_green(
    "5",
    "How should Capital Gains Tax appear on the Pro Forma?",
    "Choose one: "
    "(1) note only (words, no amount), "
    "(2) money line with an amount staff type, "
    "(3) both a note and an amount line, "
    "(4) do not show Capital Gains Tax on the Pro Forma at all.",
    "1 / 2 / 3 / 4",
)
ask_green(
    "6",
    "Is Industrial calculated the same way as Developer?",
    "Developer uses Display price + Dutiable value, stamp duty on dutiable value only, "
    "and conveyancing not on the VAT-inclusive gross. "
    "Is Industrial exactly the same, or different? "
    "If different, explain what changes (fields, stamp duty base, VAT, fees).",
    "Same as Developer / Different (explain on Others lines)",
)
ask_green(
    "7",
    "When an Industrial deal is VAT-exempt, what should change on the Pro Forma?",
    "Tell us what staff should see or hide when the property is VAT-exempt. "
    "Examples: hide VAT sentence, hide display-vs-dutiable split, change stamp duty base, "
    "change conveyancing base, or no change except a tick box that says VAT exempt.",
    "Describe the changes, or write no special change.",
)

green_para("C. Stamp duty confirmation (must answer before coding)", bold=True)
green_para(
    "Draft table again for easy reference while answering:",
    bold=False,
)
green_bullet("0 to 5,000: 1% of full value + 70. Example 4,000 -> 110")
green_bullet("Above 5,000 to 20,000: 2% of full value + 70. Example 10,000 -> 270")
green_bullet("Above 20,000 to 100,000: 3% of (value - 20,000) + 370. Example 50,000 -> 1,270")
green_bullet("Above 100,000: 4% of (value - 100,000) + 2,770. Example 500,000 -> 18,770")
green_para("")

ask_green(
    "8",
    "For property values above USD 20,000 up to USD 100,000, which stamp duty formula is correct?",
    "We need one clear rule. "
    "Option A (full value): 3% of the full property value + USD 370. "
    "Example for 50,000: (50,000 x 3%) + 370 = 1,870. "
    "Option B (excess): 3% of (value - 20,000) + USD 370. "
    "Example for 50,000: (30,000 x 3%) + 370 = 1,270. "
    "In version (3) you left Option B and the 1,270 example. "
    "Please confirm that is what T Pfigu puts on a real Pro Forma.",
    "A (full value, 1,870) / B (excess, 1,270)",
)
ask_green(
    "9",
    "For property values above USD 100,000, which stamp duty formula is correct?",
    "Option A (full value): 4% of the full property value + USD 2,770. "
    "Example for 500,000: (500,000 x 4%) + 2,770 = 22,770. "
    "Option B (excess): 4% of (value - 100,000) + USD 2,770. "
    "Example for 500,000: (400,000 x 4%) + 2,770 = 18,770. "
    "In version (3) you left Option B and the 18,770 example. Please confirm.",
    "A (full value, 22,770) / B (excess, 18,770)",
)
ask_green(
    "10",
    "For a USD 10,000 normal transfer, what TOTAL stamp duty should print on the Pro Forma?",
    "If the rule is 2% of full value + flat USD 70, then total = 200 + 70 = 270. "
    "Earlier someone said “70”; we now understand that was the flat fee only, not the total. "
    "Please confirm the number that should appear as the stamp duty line amount for 10,000.",
    "270 / other amount (write the number)",
)
ask_green(
    "11",
    "Are the four stamp duty bands and flat fees complete and correct?",
    "Please check the band edges carefully: "
    "0 to 5,000; above 5,000 to 20,000; above 20,000 to 100,000; above 100,000. "
    "Also check flat fees: 70, 70, 370, 2,770. "
    "If any edge or flat fee is wrong, write the correct bands and fees on the Others lines. "
    "Also confirm again: there is no 5% bracket.",
    "Yes, table is correct / No (write corrections on Others lines)",
)

green_para("D. Money flow and banking on the PDF", bold=True)
ask_green(
    "12",
    "When the client pays against a Pro Forma, does the whole amount go into Trust first?",
    "This is about the payment instruction on the document. "
    "If yes, the client pays one amount into the Trust bank account shown on the Pro Forma. "
    "If no, explain where each part of the money goes.",
    "Yes, whole amount to Trust first / No (explain on Others lines)",
)
ask_green(
    "13",
    "How should “conveyancing goes into our firm account” work in practice?",
    "You noted that conveyancing (about 3%, or negotiated) goes into the firm account. "
    "We need to know how that fits with Trust banking on the Pro Forma. "
    "Option A: Client still pays everything to Trust. Later, the firm moves the conveyancing "
    "portion from Trust to the firm operating account (outside or inside LexZim later). "
    "Option B: The Pro Forma shows two different bank destinations "
    "(some lines to Trust, conveyancing to firm). "
    "Option C: Something else.",
    "A / B / C. If B or C, explain which lines go where.",
)
ask_green(
    "14",
    "What should the PDF show about payment and conveyancing?",
    "Choose what the printed Pro Forma should show: "
    "(1) one Trust payment total only, "
    "(2) one Trust payment total plus a short note that conveyancing is firm income, "
    "(3) split totals by bank account, "
    "(4) something else.",
    "1 / 2 / 3 / 4",
)

green_para("E. Later conversion (optional for first build)", bold=True)
ask_green(
    "15",
    "Later, when a Pro Forma becomes final, what should staff convert it into?",
    "This does not have to be built in the first version, but we want your intended end state. "
    "Options: trust request only; fee invoice only; both; or decide case by case "
    "depending on the matter. "
    "Also say if converting should copy the lines into the new document automatically.",
    "trust request / fee invoice / both / case by case",
)

green_para("F. Short answer checklist (green)", bold=True)
green_para(
    "Quick answers here. If you need to explain more, use Others under each question above, "
    "or write under the matching number here."
)

checklist = [
    ("1", "Screen type choice: A / B / C ______"),
    ("2", "Exact type labels: ______"),
    ("3", "Extra always-on lines: none / list ______"),
    ("4", "VAT sentence: ______"),
    ("5", "Capital Gains Tax: 1 / 2 / 3 / 4 ______"),
    ("6", "Industrial same as Developer: Yes / No ______"),
    ("7", "VAT-exempt Industrial changes: ______"),
    ("8", "Stamp duty above 20,000: A (1,870) / B (1,270) ______"),
    ("9", "Stamp duty above 100,000: A (22,770) / B (18,770) ______"),
    ("10", "USD 10,000 stamp duty total: 270 / other ______"),
    ("11", "Bands and flat fees correct: Yes / No ______"),
    ("12", "Whole payment to Trust first: Yes / No ______"),
    ("13", "Conveyancing to firm: A / B / C ______"),
    ("14", "PDF payment display: 1 / 2 / 3 / 4 ______"),
    ("15", "Later convert path: ______"),
]
for num, text in checklist:
    p = doc.add_paragraph()
    r = p.add_run(f"Q{num}. {text}")
    r.font.color.rgb = GREEN
    r.bold = True
    r.font.size = Pt(11)
    blank_line("Others / further explanation:", color=GREEN)
    para("")

green_para(
    "Thank you. Once this green section is answered, we can lock the stamp duty engine "
    "and money-flow rules and start building.",
    bold=True,
)
green_para("End of green questions section.")

doc.save(path)
print(f"Saved {path}")
